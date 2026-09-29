"""셀러 콘텐츠 캘린더 앱 빌더. (셀러별 설정은 config.json)

구글 시트(공개 xlsx export)를 내려받아 캘린더/제품/월별 데이터를 JSON으로 정리한 뒤
template.html 에 넣어 모바일용 단일 HTML(duckduck_calendar.html)을 만든다.

    python build_app.py            # 시트 새로 받아서 빌드
    python build_app.py --local    # 이미 받아둔 sheet.xlsx 로 빌드
"""
import datetime as dt
import json
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl

import sheet_media

HERE = Path(__file__).parent
CONFIG = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
SHEET_ID = CONFIG["sheet_id"]
YEAR = CONFIG.get("year", 2026)
XLSX = HERE / "sheet.xlsx"
TEMPLATE = HERE / "template.html"
OUT = HERE / "duckduck_calendar.html"
MEDIA = HERE / "media"  # 시트 이미지를 줄여서 저장하는 곳 (앱에서는 media/파일명)
IMAGES = {}  # 시트 이름 → {(행, 열): [파일명]}
FORMATS = {}  # (시트 이름, 행, 열) → 서식·링크 정보 (sheet_media.CellFormat)

CAL_COLS = "BCDEFGH"  # 월~일
MEETING_FILL = "FFCFE2F3"
PALETTE = ["#A4C2F4", "#F9CB9C", "#B6D7A8", "#D9D2E9", "#FFE599", "#EA9999", "#A2C4C9", "#D5A6BD"]
HOLIDAY_FILL = "FFF3F3F3"


def download():
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    with urllib.request.urlopen(url) as r:
        XLSX.write_bytes(r.read())


def norm(s):
    return re.sub(r"[\s/()]|\d차", "", s or "")


def text(v):
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def fill(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


def iso(d):
    return d.strftime("%Y-%m-%d")


def shown(ws, r, c, plain):
    """화면에 보일 칸 글자: 시트의 굵게·색·크기·링크 표시를 넣어서 돌려준다 (구조 판단은 plain 으로)"""
    f = FORMATS.get((ws.title.strip(), r, c))
    if not f or not plain:
        return plain
    lead = len(f.text) - len(f.text.lstrip())
    if f.text.strip() == plain:
        return f.render(lead, lead + len(plain))
    return f.render_text(plain)


def notes_of(ws, r, cells, min_len=40):
    """행의 긴 글(목표·참고사항 등)을 서식 포함해서"""
    return [shown(ws, r, c, v) for c, v in cells if len(v) > min_len]


def row_images(ws, r, cols=None):
    """r 행에 걸린 이미지 (열 순서)"""
    cells = IMAGES.get(ws.title.strip(), {})
    return [f"media/{n}" for c in sorted(c for (rr, c) in cells if rr == r and (cols is None or c in cols))
            for n in cells[(r, c)]]


def top_left_only(ws):
    """merged range 의 좌상단이 아닌 셀 좌표 집합 (값 없음 취급)"""
    return {(r, c) for m in ws.merged_cells.ranges
            for r in range(m.min_row, m.max_row + 1)
            for c in range(m.min_col, m.max_col + 1)
            if (r, c) != (m.min_row, m.min_col)}


# ---------------------------------------------------------------- 캘린더 탭
def parse_calendar(ws, product_keys):
    legend, events = {}, []
    for row in ws.iter_rows(min_col=7, max_col=8):
        g, h = row
        if not text(g.value) and fill(g) and norm(text(h.value)) in product_keys:
            legend[fill(g)] = text(h.value)

    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    month, week_rows, header_rows = None, [], set()
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        month_label = text(ws.cell(r, 12).value)
        if re.fullmatch(r"\d{1,2}월", month_label):
            month = int(month_label[:-1])
            header_rows.add(r)
            continue
        dates = {}
        for i, c in enumerate(CAL_COLS):
            v = ws[f"{c}{r}"].value
            if isinstance(v, dt.datetime) and v.year == YEAR:
                dates[i] = v.day
            elif isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and 1 <= v <= 31:
                dates[i] = int(v)
        if len(dates) == 1 and not any(isinstance(ws[f"{c}{r}"].value, dt.datetime) for c in CAL_COLS):
            dates = {}  # 숫자 하나만 있는 행은 날짜 행으로 보지 않음
        if dates and month:
            week_rows.append((r, month, dates))

    for idx, (r, month, dates) in enumerate(week_rows):
        nxt = week_rows[idx + 1][0] if idx + 1 < len(week_rows) else ws.max_row + 1
        for rr in range(r + 1, min(nxt, r + 7)):
            if rr in header_rows:
                break
            for i, c in enumerate(CAL_COLS):
                cell = ws[f"{c}{rr}"]
                t = text(cell.value)
                if not t or isinstance(cell.value, dt.datetime) or i not in dates:
                    continue
                # 셀에 적힌 날짜가 월 헤더와 다른 경우(5월 블록 오류)도 월 헤더 기준으로 보정
                start = dt.date(YEAR, month, dates[i])
                span = 1
                m = merged.get((rr, cell.column))
                if m:
                    span = m.max_col - m.min_col + 1
                f = fill(cell)
                kind = "meeting" if f == MEETING_FILL else "holiday" if f == HOLIDAY_FILL else "content"
                events.append({
                    "date": iso(start),
                    "end": iso(start + dt.timedelta(days=span - 1)),
                    "text": t,
                    "product": legend.get(f),
                    "kind": kind,
                    "order": rr - r,
                })
    colors = {name: "#" + rgb[2:] for rgb, name in legend.items()}
    return events, colors


# ---------------------------------------------------------------- 제품 탭
DATE_RE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})")
FORMAT_RE = re.compile(r"(스토리|릴스|게시글|게시물|캐러셀|피드|무물|라이브|라방)")


def parse_period(s):
    m = re.search(r"(\d{1,2})/(\d{1,2})\s*\([^)]*\)\s*~\s*(\d{1,2})/(\d{1,2})", s)
    if not m:
        return None
    a = dt.date(YEAR, int(m[1]), int(m[2]))
    b = dt.date(YEAR, int(m[3]), int(m[4]))
    return [iso(a), iso(b)]


def split_topic(b):
    lines = [l.strip() for l in b.split("\n") if l.strip()]
    if not lines:
        return "", "", []
    fmt = lines[0]
    m = re.match(r"\[(.+?)\]\s*(.*)", fmt)
    if m:
        fmt = m[1].strip()
        if m[2].strip():
            lines = [fmt, m[2].strip(), *lines[1:]]
    if not FORMAT_RE.match(fmt):
        return "기타", " ".join(lines), []
    if fmt.startswith("스토리"):
        rest = " ".join(l.lstrip("*").strip() for l in lines[1:])
        return "스토리", "", [rest] if rest else []
    title = [l for l in lines[1:] if not l.startswith("*")]
    notes = [l.lstrip("*").strip() for l in lines[1:] if l.startswith("*")]
    return fmt, " ".join(title), notes


def parse_product(ws, colors):
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    header_row = next(r for r in range(1, 15) if val(r, 1) == "일정")
    headers = {c: val(header_row, c) for c in range(3, ws.max_column + 1) if val(header_row, c)}

    info, period, product_line = [], None, None
    for r in range(1, header_row):
        t = val(r, 1)
        if not t:
            continue
        for block in re.split(r"\n\s*\n(?=📍)", t):
            block = block.strip()
            head, _, body = block.partition("\n")
            head = head.replace("📍", "").strip()
            if "공구일정" in head:
                period = parse_period(head)
            if "공구상품" in head:
                product_line = head.split(":", 1)[-1].strip()
            if ":" in head and not body:
                k, v = head.split(":", 1)
                info.append({"label": k.strip(), "text": shown(ws, r, 1, v.strip())})
            else:
                label, _, rest = head.partition(":")
                body = (rest.strip() + "\n" + body).strip() if rest.strip() else body
                info.append({"label": label.strip(), "text": shown(ws, r, 1, body.strip())})

    # A/B 는 병합 셀이 많아 값 전파
    a_merge, b_merge = {}, {}
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            if m.min_col == 1:
                a_merge[r] = m.min_row
            if m.min_col == 2:
                b_merge[r] = m.min_row

    info_images = [i for r in range(1, header_row) for i in row_images(ws, r)]
    pending = row_images(ws, header_row)  # 헤더 줄에 걸쳐 놓인 이미지는 첫 행 것
    slots, cur, last_a = [], None, ""
    for r in range(header_row + 1, ws.max_row + 1):
        a = text(ws.cell(a_merge.get(r, r), 1).value)
        b = text(ws.cell(b_merge.get(r, r), 2).value)
        fields = {}
        for c, h in headers.items():
            v = val(r, c)
            if v and v != "-":
                fields[h] = shown(ws, r, c, v)
        imgs = pending + row_images(ws, r)
        pending = []
        if not (a or b or fields or imgs):
            continue
        new_a = a and (r not in a_merge or a_merge[r] == r)
        new_b = b and (r not in b_merge or b_merge[r] == r)
        if cur is None or new_a or new_b:
            a = a or last_a
            last_a = a
            m = DATE_RE.search(a)
            fmt, title, notes = split_topic(b)
            cur = {
                "label": a.split("\n")[0].split(" ")[0].strip(),
                "date": iso(dt.date(YEAR, int(m[1]), int(m[2]))) if m else None,
                "format": fmt or "기타",
                "title": title,
                "notes": notes,
                "items": [],
            }
            slots.append(cur)
        if fields:
            if cur["format"].startswith("스토리") or not cur["items"]:
                cur["items"].append(fields)
            else:
                last = cur["items"][-1]
                for k, v in fields.items():
                    last[k] = (last[k] + "\n\n" + v) if k in last else v
        if imgs:
            if not cur["items"]:
                cur["items"].append({})
            cur["items"][-1].setdefault("__images", []).extend(imgs)

    name = ws.title.strip()
    base = norm(name)
    color = next((c for n, c in colors.items() if norm(n) == base), None)
    if color is None:
        color = PALETTE[sum(map(ord, base)) % len(PALETTE)]
    round_m = re.search(r"(\d)차", name)
    return {
        "id": "p" + re.sub(r"\W", "", base) + (round_m[1] if round_m else ""),
        "name": re.sub(r"\s*\(?\d차\)?", "", name).strip(),
        "round": f"{round_m[1]}차" if round_m else "",
        "fullName": product_line,
        "color": color,
        "period": period,
        "info": info,
        "images": info_images,
        "slots": slots,
    }


# ---------------------------------------------------------------- 월별 탭
GRID_LABELS = {"스토리", "주제", "기획 의도", "팔로워 반응", "콘텐츠 구성", "콘텐츠 플로우", "캡션 참고",
               "콘텐츠", "기대 효과", "기획 의도 & 기대 효과", "팔로워 예상 반응", "피드 비주얼", "피드 참고"}
DAY_RE = re.compile(r"^(\d{1,2})(?:\.0)?(?:\s*\((.+)\))?$")


def parse_month_grid(ws, month):
    """3~5월: 주 단위 그리드 (날짜 행 + 라벨 행)"""
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    days, notes, extra = {}, [], []
    week, first_week, in_extra = None, None, False
    for r in range(1, ws.max_row + 1):
        cells = {c: val(r, c) for c in range(2, 9)}
        a = val(r, 1)
        day_hits = {c: DAY_RE.match(v) for c, v in cells.items() if v}
        day_hits = {c: m for c, m in day_hits.items() if m}
        if len(day_hits) >= 1 and (a == "날짜" or len(day_hits) >= 3 or week is None and len(day_hits) >= 1):
            week = {}
            for c, m in day_hits.items():
                d = int(m[1])
                key = iso(dt.date(YEAR, month, d))
                week[c] = key
                days.setdefault(key, {"holiday": m[2], "fields": []})
            if first_week is None:
                first_week = r
            for c, v in cells.items():
                if len(v) > 40 and c not in day_hits:
                    notes.append(shown(ws, r, c, v))
            continue
        if week is None:
            notes += notes_of(ws, r, [(1, a), *cells.items()])
            continue
        label = a.replace("\n", " ")
        if label == "날짜":
            continue
        for c, key in week.items():
            imgs = row_images(ws, r, {c})
            if imgs:
                days[key].setdefault("images", []).extend(imgs)
        if label:
            in_extra = label not in GRID_LABELS
        for c, v in cells.items():
            if len(v) < 2:
                continue
            if in_extra:
                if label:
                    extra.append({"label": label, "text": shown(ws, r, c, v)})
                    label = ""
                else:
                    extra[-1]["text"] += "\n\n" + shown(ws, r, c, v)
            elif c in week and label:
                days[week[c]]["fields"].append({"label": label, "text": shown(ws, r, c, v)})
            elif len(v) > 40:
                notes.append(shown(ws, r, c, v))
    day_list = [{"date": k, **v} for k, v in sorted(days.items()) if v["fields"] or v["holiday"] or v.get("images")]
    return {"kind": "grid", "notes": notes, "days": day_list, "extra": extra}


def parse_month_sections(ws):
    """6~8월: 카테고리 섹션별 기획 리스트"""
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    notes, sections, cur, item = [], [], None, None
    cols = {2: "기획 의도", 3: "콘텐츠 주제", 4: "콘텐츠 플로우", 5: "캡션 참고", 6: "비고"}
    for r in range(2, ws.max_row + 1):
        raw_a = ws.cell(r, 1).value if (r, 1) not in skip else None
        a = text(raw_a)
        rest = {c: val(r, c) for c in range(2, 7)}
        if a.startswith("포인트"):
            continue
        if rest.get(2) == "팔로워 반응" and not a:
            continue
        imgs = row_images(ws, r)
        if item and imgs and not isinstance(raw_a, bool) and a not in ("True", "False"):
            item.setdefault("images", []).extend(imgs)
        if isinstance(raw_a, bool) or a in ("True", "False"):
            item = {"done": a == "True" or raw_a is True, "fields": {}}
            if imgs:
                item["images"] = imgs
            for c, v in rest.items():
                if v:
                    item["fields"][cols[c]] = v if c == 3 else shown(ws, r, c, v)
            item["title"] = item["fields"].pop("콘텐츠 주제", "")
            if "[" not in item["title"][:12]:  # 대괄호 포맷이 없는 건 스토리 아이디어
                if item["title"]:
                    item["fields"] = {"스토리 내용": item["title"], **item["fields"]}
                head = item["title"].split(" - ")[0] if " - " in item["title"][:20] else ""
                item["title"] = head or "스토리 아이디어"
                item["story"] = True
            if cur is None:
                cur = {"title": "", "items": []}
                sections.append(cur)
            cur["items"].append(item)
            continue
        if a and not any(rest.values()) and r > 3:
            title, _, sub = a.partition("*")
            cur = {"title": title.strip(), "sub": sub.strip(), "items": []}
            sections.append(cur)
            item = None
            continue
        if not sections:
            notes += notes_of(ws, r, [(1, a), *rest.items()])
            continue
        if item and not a:
            if rest.get(2):
                item["reaction"] = (item.get("reaction", "") + "\n" + shown(ws, r, 2, rest[2])).strip()
            for c in range(3, 7):
                if rest.get(c):
                    k = cols[c]
                    item["fields"][k] = (item["fields"].get(k, "") + "\n\n" + shown(ws, r, c, rest[c])).strip()
    return {"kind": "sections", "notes": notes, "sections": sections}


def parse_month_weeks(ws):
    """주차별 섹션: 'n주차 콘텐츠' 제목 → '유형' 헤더 → 항목 행 + 이어지는 보조 행"""
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    notes, sections, headers, item = [], [], {}, None
    for r in range(1, ws.max_row + 1):
        a = val(r, 1)
        row = {c: val(r, c) for c in range(2, ws.max_column + 1)}
        if "주차" in a:
            title, _, rng = a.partition("(")
            sections.append({"title": title.replace("콘텐츠", "").strip(),
                             "sub": rng.rstrip(")").strip(), "items": []})
            item = None
            continue
        if a == "유형":
            headers = {c: v for c, v in row.items() if v}
            continue
        if not sections:
            notes += notes_of(ws, r, [(1, a), *row.items()])
            continue
        done = any(ws.cell(r, c).value is True for c in range(2, ws.max_column + 1))
        imgs = row_images(ws, r)
        if item and imgs and not a:
            item.setdefault("images", []).extend(imgs)
        if a:
            fields = {headers[c]: (v if c == 2 else shown(ws, r, c, v)) for c, v in row.items() if v and c in headers}
            topic = fields.pop(headers.get(2, ""), "")
            item = {"done": done, "title": f"[{a}] {topic}".strip(), "fields": fields}
            if imgs:
                item["images"] = imgs
            sections[-1]["items"].append(item)
        elif item:
            for c, v in row.items():
                if not v or c not in headers:
                    continue
                if "팔로워" in headers[c] and headers[c] == headers.get(3):
                    item["reaction"] = (item.get("reaction", "") + "\n" + shown(ws, r, c, v)).strip()
                else:
                    k = headers[c]
                    item["fields"][k] = (item["fields"].get(k, "") + "\n\n" + shown(ws, r, c, v)).strip()
    return {"kind": "sections", "notes": notes, "sections": sections}


def parse_month(ws):
    m = re.match(r"(\d{2})\.(\d{2})", ws.title.strip())
    month = int(m[2])
    has_grid = any(text(ws.cell(r, 1).value) == "날짜" for r in range(1, 10))
    has_weeks = any(text(ws.cell(r, 1).value) == "유형" for r in range(1, ws.max_row + 1))
    data = (parse_month_grid(ws, month) if has_grid
            else parse_month_weeks(ws) if has_weeks else parse_month_sections(ws))
    data.update({"id": f"m{month:02d}", "month": month, "name": f"{month}월"})
    return data


# ---------------------------------------------------------------- main
def main():
    if "--local" not in sys.argv:
        download()
    wb = openpyxl.load_workbook(XLSX, data_only=True, rich_text=True)
    try:
        links = sheet_media.fetch_links(SHEET_ID)
    except Exception as e:  # 시트 HTML 보기를 못 읽으면 xlsx 의 셀 링크만 사용
        print("links: htmlview 실패, xlsx 링크로 대체 -", e)
        links = {}
    FORMATS.update(sheet_media.collect_formats(wb, links))
    IMAGES.update(sheet_media.extract_images(XLSX, MEDIA))
    print("images:", sum(len(v) for cells in IMAGES.values() for v in cells.values()))
    sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    cal = next(ws for ws in sheets if "캘린더" in ws.title)
    product_sheets = [ws for ws in sheets
                      if any(text(ws.cell(r, 1).value) == "일정" for r in range(1, 15))]
    month_sheets = [ws for ws in sheets if re.match(r"\d{2}\.\d{2}", ws.title.strip())]
    events, colors = parse_calendar(cal, {norm(ws.title) for ws in product_sheets})
    products = [parse_product(ws, colors) for ws in product_sheets]
    months = [parse_month(ws) for ws in month_sheets]
    for e in events:  # 범례 이름 → 제품 탭 id (같은 제품 여러 차수면 날짜가 가까운 차수)
        cands = [p for p in products if e["product"] and norm(p["name"]) == norm(e["product"])]

        def dist(p):
            ds = [s["date"] for s in p["slots"] if s["date"]] + (p["period"] or [])
            return min(abs((dt.date.fromisoformat(d) - dt.date.fromisoformat(e["date"])).days) for d in ds)
        e["product"] = min(cands, key=dist)["id"] if cands else None
    products.sort(key=lambda p: (p["period"] or ["9999"])[0], reverse=True)
    months.sort(key=lambda m: m["month"], reverse=True)
    data = {
        "updated": dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).strftime("%Y-%m-%d %H:%M"),
        "sheetUrl": f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit",
        "year": YEAR,
        "events": events,
        "products": products,
        "months": months,
    }
    (HERE / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    if TEMPLATE.exists():
        html = TEMPLATE.read_text(encoding="utf-8")
        payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
        html = html.replace("__APP_NAME__", CONFIG["app_name"]).replace("__SELLER__", CONFIG["seller"])
        OUT.write_text(html.replace("/*__DATA__*/null", payload), encoding="utf-8")
        print("built", OUT)
    print(f"events={len(events)} products={len(products)} months={len(months)}")


if __name__ == "__main__":
    main()
